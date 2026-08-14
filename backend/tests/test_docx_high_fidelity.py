import base64
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from app.translate import to_translate, word
from app.translate.docx_inline import build_inline_plan


PNG_1X1 = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII='
)


def _styled_translation(plan, values):
    parts = []
    value_iter = iter(values)
    for kind, item in plan.sequence:
        if kind == 'anchor':
            parts.append(item.token)
        else:
            parts.extend((item.start_token, next(value_iter), item.end_token))
    return ''.join(parts)


def _add_hyperlink(paragraph, text, url):
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relation_id)
    run = OxmlElement('w:r')
    run_properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    run_properties.append(color)
    run.append(run_properties)
    text_element = OxmlElement('w:t')
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return relation_id


def _add_page_field(paragraph):
    begin_run = paragraph.add_run()._r
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin_run.append(begin)

    instruction_run = paragraph.add_run()._r
    instruction = OxmlElement('w:instrText')
    instruction.text = ' PAGE '
    instruction_run.append(instruction)

    separate_run = paragraph.add_run()._r
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    separate_run.append(separate)
    paragraph.add_run('1')

    end_run = paragraph.add_run()._r
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    end_run.append(end)


class DocxHighFidelityTests(unittest.TestCase):
    def test_mixed_run_styles_survive_inline_replacement(self):
        document = Document()
        paragraph = document.add_paragraph()
        bold = paragraph.add_run('Hello ')
        bold.bold = True
        italic = paragraph.add_run('world')
        italic.italic = True

        plan = build_inline_plan(paragraph)
        self.assertEqual(len(plan.spans), 2)
        plan.apply(_styled_translation(plan, ['Bonjour ', 'monde']))

        self.assertEqual(paragraph.text, 'Bonjour monde')
        self.assertTrue(paragraph.runs[0].bold)
        self.assertTrue(paragraph.runs[1].italic)

    def test_image_anchor_keeps_text_on_both_sides(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / 'pixel.png'
            image_path.write_bytes(PNG_1X1)
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run('Before ')
            paragraph.add_run().add_picture(str(image_path))
            paragraph.add_run(' after')

            plan = build_inline_plan(paragraph)
            drawing_anchor = next(
                anchor for anchor in plan.anchors if anchor.kind == 'drawing'
            )
            translated = _styled_translation(plan, ['Avant ', ' apres'])
            self.assertIn(drawing_anchor.token, translated)
            plan.apply(translated)

            children = list(paragraph._p)
            drawing_index = next(
                index for index, child in enumerate(children)
                if child.findall('.//' + qn('w:drawing'))
            )
            before_text = ''.join(
                node.text or ''
                for child in children[:drawing_index]
                for node in child.findall('.//' + qn('w:t'))
            )
            after_text = ''.join(
                node.text or ''
                for child in children[drawing_index + 1:]
                for node in child.findall('.//' + qn('w:t'))
            )
            self.assertEqual(before_text, 'Avant ')
            self.assertEqual(after_text, ' apres')
            output_path = Path(temp_dir) / 'image-output.docx'
            document.save(output_path)
            reopened = Document(output_path)
            self.assertEqual(len(reopened.inline_shapes), 1)
            self.assertEqual(reopened.paragraphs[0].text, 'Avant  apres')

    def test_hyperlink_relationship_and_visible_text_survive(self):
        document = Document()
        paragraph = document.add_paragraph('Visit ')
        relation_id = _add_hyperlink(paragraph, 'the docs', 'https://example.com')
        paragraph.add_run(' now')

        plan = build_inline_plan(paragraph)
        plan.apply(_styled_translation(plan, ['Consultez ', 'la documentation', ' maintenant']))

        hyperlink = paragraph._p.find(qn('w:hyperlink'))
        self.assertIsNotNone(hyperlink)
        self.assertEqual(hyperlink.get(qn('r:id')), relation_id)
        self.assertEqual(
            ''.join(node.text or '' for node in hyperlink.findall('.//' + qn('w:t'))),
            'la documentation',
        )
        self.assertEqual(paragraph.part.rels[relation_id].target_ref, 'https://example.com')
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / 'hyperlink-output.docx'
            document.save(output_path)
            reopened = Document(output_path)
            reopened_hyperlink = reopened.paragraphs[0]._p.find(qn('w:hyperlink'))
            reopened_relation_id = reopened_hyperlink.get(qn('r:id'))
            self.assertEqual(
                reopened.paragraphs[0].part.rels[reopened_relation_id].target_ref,
                'https://example.com',
            )

    def test_field_and_bookmark_nodes_remain_unchanged(self):
        document = Document()
        paragraph = document.add_paragraph('Page ')
        _add_page_field(paragraph)
        paragraph.add_run(' in report')
        bookmark_start = OxmlElement('w:bookmarkStart')
        bookmark_start.set(qn('w:id'), '7')
        bookmark_start.set(qn('w:name'), 'report_end')
        bookmark_end = OxmlElement('w:bookmarkEnd')
        bookmark_end.set(qn('w:id'), '7')
        paragraph._p.insert(len(paragraph._p) - 1, bookmark_start)
        paragraph._p.append(bookmark_end)

        field_xml = [
            etree.tostring(node)
            for node in paragraph._p.findall('.//' + qn('w:fldChar'))
        ]
        plan = build_inline_plan(paragraph)
        plan.apply(_styled_translation(plan, ['Page ', ' in report']))

        self.assertEqual(
            [
                etree.tostring(node)
                for node in paragraph._p.findall('.//' + qn('w:fldChar'))
            ],
            field_xml,
        )
        self.assertEqual(
            paragraph._p.find(qn('w:bookmarkStart')).get(qn('w:id')), '7'
        )
        self.assertEqual(
            paragraph._p.find(qn('w:bookmarkEnd')).get(qn('w:id')), '7'
        )
        self.assertEqual(
            ''.join(node.text or '' for node in paragraph._p.findall('.//' + qn('w:instrText'))),
            ' PAGE ',
        )

    @patch('app.translate.to_translate.time.sleep')
    @patch('app.translate.to_translate._translate_openai')
    def test_style_token_failure_uses_main_style_fallback(self, translate_openai, sleep):
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run('Hello ').bold = True
        paragraph.add_run('world').italic = True
        plan = build_inline_plan(paragraph)
        item = {'text': plan.payload, **plan.translation_metadata()}
        translate_openai.side_effect = [
            'broken style tokens',
            'broken style tokens',
            'broken style tokens',
            'Bonjour le monde',
        ]

        result = to_translate._translate_text_block(
            {'id': 1, 'server': 'openai', 'model': 'model'}, item
        )

        self.assertTrue(result['fallback_used'])
        plan.apply(result['translated_text'], fallback_used=True)
        self.assertEqual(paragraph.text, 'Bonjour le monde')
        self.assertTrue(paragraph.runs[0].bold)
        self.assertEqual(paragraph.runs[1].text, '')

    def test_word_start_preserves_table_paragraphs_and_reopens_output(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'source.docx'
            target = Path(temp_dir) / 'target.docx'
            document = Document()
            document.add_paragraph('Body text')
            table = document.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            cell.paragraphs[0].add_run('First paragraph')
            cell.add_paragraph('Second paragraph')
            document.sections[0].header.paragraphs[0].text = 'Header text'
            document.save(source)

            def fake_translate_batch(trans, texts, event):
                replacements = {
                    'Body text': 'Corps',
                    'First paragraph': 'Premier',
                    'Second paragraph': 'Deuxieme',
                    'Header text': 'En-tete',
                }
                for item in texts:
                    translated = item['text']
                    for original, replacement in replacements.items():
                        translated = translated.replace(original, replacement)
                    item['text'] = translated
                    item['count'] = 1
                    item['complete'] = True
                return True

            with patch.object(word.to_translate, 'translate_batch', fake_translate_batch), patch.object(
                word.to_translate, 'complete', return_value=True
            ):
                success = word.start({
                    'id': 1,
                    'file_path': str(source),
                    'target_file': str(target),
                    'type': 'trans_all_only_inherit',
                    'lang': '法语',
                })

            self.assertTrue(success)
            reopened = Document(target)
            self.assertEqual(reopened.paragraphs[0].text, 'Corps')
            self.assertEqual(reopened.tables[0].cell(0, 0).paragraphs[0].text, 'Premier')
            self.assertEqual(reopened.tables[0].cell(0, 0).paragraphs[1].text, 'Deuxieme')
            self.assertEqual(reopened.sections[0].header.paragraphs[0].text, 'En-tete')

    def test_bilingual_mode_appends_plain_translation_without_tokens(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / 'source.docx'
            target = Path(temp_dir) / 'target.docx'
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run('Hello ').bold = True
            paragraph.add_run('world').italic = True
            document.save(source)

            def fake_translate_batch(trans, texts, event):
                for item in texts:
                    item['text'] = item['text'].replace('Hello ', 'Bonjour ').replace(
                        'world', 'monde'
                    )
                    item['count'] = 2
                    item['complete'] = True
                    item['fallback_used'] = False
                return True

            with patch.object(word.to_translate, 'translate_batch', fake_translate_batch), patch.object(
                word.to_translate, 'complete', return_value=True
            ):
                success = word.start({
                    'id': 2,
                    'file_path': str(source),
                    'target_file': str(target),
                    'type': 'trans_all_both_inherit',
                    'lang': '法语',
                })

            self.assertTrue(success)
            output_text = Document(target).paragraphs[0].text
            self.assertEqual(output_text, 'Hello world\nBonjour monde')
            self.assertNotIn('⟦DOCX_', output_text)

    def test_invalid_saved_document_is_removed(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / 'broken.docx'

            def write_broken(path):
                Path(path).write_bytes(b'not a docx')

            fake_document = SimpleNamespace(save=write_broken)
            with patch('app.translate.word.Document', side_effect=ValueError('broken')):
                with self.assertRaisesRegex(ValueError, 'broken'):
                    word._save_and_validate_document(fake_document, str(target))
            self.assertFalse(target.exists())


if __name__ == '__main__':
    unittest.main()
