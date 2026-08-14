import base64
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

from app.translate import powerpoint, to_translate
from app.translate.pptx_inline import build_inline_plan


PNG_1X1 = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9Zl1sAAAAASUVORK5CYII='
)


def _styled_translation(plan, values):
    parts = []
    value_iter = iter(values)
    for kind, item in plan.sequence:
        if kind == 'anchor':
            parts.append(item.token)
        else:
            parts.extend((item.start_token, next(value_iter), item.end_token))
    return ''.join(parts)


def _all_shapes(slide):
    for shape in slide.shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shape.shapes


def _all_text(slide):
    values = []
    for shape in _all_shapes(slide):
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    values.extend(p.text for p in cell.text_frame.paragraphs)
        elif shape.has_text_frame:
            values.extend(p.text for p in shape.text_frame.paragraphs)
    return values


def _fake_translate_batch(trans, texts, event):
    replacements = {
        'Hello ': 'Bonjour ',
        'world': 'monde',
        'Documentation': 'Documentation francaise',
        'Second shape': 'Deuxieme forme avec un texte beaucoup plus long mais complet',
        'Grouped text': 'Texte groupe',
        'First paragraph': 'Premier paragraphe',
        'Second paragraph': 'Deuxieme paragraphe',
    }
    for item in texts:
        translated = item['text']
        for original, replacement in replacements.items():
            translated = translated.replace(original, replacement)
        item['text'] = translated
        item['count'] = 1
        item['complete'] = True
        item['fallback_used'] = False
    return True


def _build_complex_presentation(path: Path, image_path: Path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    first = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.8))
    paragraph = first.text_frame.paragraphs[0]
    paragraph.clear()
    bold = paragraph.add_run()
    bold.text = 'Hello '
    bold.font.bold = True
    bold.font.size = Pt(20)
    italic = paragraph.add_run()
    italic.text = 'world'
    italic.font.italic = True
    italic.font.size = Pt(18)

    link_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(3), Inches(0.6))
    link_run = link_box.text_frame.paragraphs[0].add_run()
    link_run.text = 'Documentation'
    link_run.hyperlink.address = 'https://example.com/docs'

    second = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(3), Inches(0.8))
    second.text_frame.paragraphs[0].text = 'Second shape'
    second.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    second.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT

    group = slide.shapes.add_group_shape()
    grouped = group.shapes.add_textbox(Inches(4), Inches(0.5), Inches(2), Inches(0.8))
    grouped.text_frame.paragraphs[0].text = 'Grouped text'

    table_shape = slide.shapes.add_table(1, 1, Inches(0.5), Inches(2.2), Inches(3), Inches(1.4))
    cell = table_shape.table.cell(0, 0)
    cell.text_frame.paragraphs[0].text = 'First paragraph'
    cell.text_frame.add_paragraph().text = 'Second paragraph'

    slide.shapes.add_picture(str(image_path), Inches(4), Inches(2.2), Inches(0.5), Inches(0.5))
    chart_data = ChartData()
    chart_data.categories = ['A', 'B']
    chart_data.add_series('Series', (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(4), Inches(3), Inches(3), Inches(2), chart_data,
    )
    presentation.save(path)


class PptxInlineTests(unittest.TestCase):
    def test_openai_prompt_identifies_pptx_protected_markers(self):
        client = MagicMock()
        choice = MagicMock()
        choice.message.content = 'translated'
        client.chat.completions.create.return_value.choices = [choice]
        result = to_translate._translate_openai(
            {
                'id': 3,
                'lang': '英语',
                'prompt': 'Translate to {target_lang}',
                'extension': '.pptx',
                '_openai_client': client,
            },
            '⟦PPTX_STYLE_START_1⟧Hello⟦PPTX_STYLE_END_1⟧',
            'model-a',
        )
        self.assertEqual(result, 'translated')
        prompt = client.chat.completions.create.call_args.kwargs['messages'][0]['content']
        self.assertIn('PPTX', prompt)
        self.assertIn('不可变结构标记', prompt)

    def test_mixed_run_styles_survive_inline_replacement(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        paragraph.clear()
        bold = paragraph.add_run()
        bold.text = 'Hello '
        bold.font.bold = True
        italic = paragraph.add_run()
        italic.text = 'world'
        italic.font.italic = True

        plan = build_inline_plan(paragraph)
        self.assertEqual(len(plan.spans), 2)
        plan.apply(_styled_translation(plan, ['Bonjour ', 'monde']))

        self.assertEqual(paragraph.text, 'Bonjour monde')
        self.assertTrue(paragraph.runs[0].font.bold)
        self.assertTrue(paragraph.runs[1].font.italic)

    def test_hyperlink_field_break_and_tab_are_structural_anchors(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(5), Inches(1)).text_frame.paragraphs[0]
        paragraph.clear()
        paragraph.add_run().text = 'Visit '
        linked = paragraph.add_run()
        linked.text = 'docs'
        linked.hyperlink.address = 'https://example.com'
        paragraph.add_line_break()
        paragraph.add_run().text = 'Page '
        field = parse_xml(
            '<a:fld %s id="{11111111-1111-1111-1111-111111111111}" type="slidenum">'
            '<a:rPr/><a:t>1</a:t></a:fld>' % nsdecls('a')
        )
        paragraph._p.insert(len(paragraph._p) - 1, field)
        paragraph.add_run().text = '\tend'

        relation_target = linked.hyperlink.address
        plan = build_inline_plan(paragraph)
        translated = plan.payload.replace('Visit ', 'Consultez ').replace(
            'docs', 'documentation'
        ).replace('Page ', 'Page ').replace('end', 'fin')
        plan.apply(translated)

        self.assertEqual(linked.hyperlink.address, relation_target)
        self.assertEqual(linked.text, 'documentation')
        self.assertIn('\v', paragraph.text)
        self.assertIn('\tfin', paragraph.text)
        self.assertEqual(paragraph._p.find(qn('a:fld')).text, '1')

    def test_cjk_font_node_precedes_hyperlink_node(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        paragraph.clear()
        linked = paragraph.add_run()
        linked.text = 'Documentation'
        linked.hyperlink.address = 'https://example.com'
        plan = build_inline_plan(paragraph)

        powerpoint._ensure_plan_font_compatibility(plan, '中文')

        run_properties = linked._r.get_or_add_rPr()
        child_tags = [child.tag for child in run_properties]
        self.assertLess(child_tags.index(qn('a:latin')), child_tags.index(qn('a:ea')))
        self.assertLess(child_tags.index(qn('a:ea')), child_tags.index(qn('a:hlinkClick')))
        self.assertEqual(linked.hyperlink.address, 'https://example.com')

    def test_inherited_font_uses_text_frame_autofit(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        placeholder = slide.placeholders[1]
        paragraph = placeholder.text_frame.paragraphs[0]
        paragraph.text = 'Short placeholder'
        self.assertIsNone(paragraph.runs[0].font.size)
        geometry = (
            placeholder.left, placeholder.top, placeholder.width, placeholder.height
        )
        blocks = powerpoint._extract_all_blocks(presentation)
        block = next(item for item in blocks if item.shape_id == placeholder.shape_id)
        block.translated_text = block.inline_plan.payload.replace(
            'Short placeholder',
            'A substantially longer translated placeholder sentence that must remain complete',
        )
        block.complete = True

        powerpoint._apply_translation_mode(presentation, blocks, '英语')

        self.assertEqual(
            (placeholder.left, placeholder.top, placeholder.width, placeholder.height),
            geometry,
        )
        self.assertIsNone(paragraph.runs[0].font.size)
        self.assertEqual(placeholder.text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)
        scale = placeholder.text_frame._txBody.bodyPr.normAutofit.fontScale
        self.assertGreaterEqual(scale, 60.0)
        self.assertLess(scale, 100.0)

    def test_existing_text_to_fit_scale_is_never_enlarged(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame
        text_frame.paragraphs[0].text = 'Short text'
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        text_frame._txBody.bodyPr.normAutofit.fontScale = 50.0
        blocks = powerpoint._extract_all_blocks(presentation)
        block = blocks[0]
        block.translated_text = block.inline_plan.payload.replace(
            'Short text', 'A longer translated sentence'
        )

        powerpoint._apply_translation_mode(presentation, blocks, '英语')

        self.assertEqual(text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)
        self.assertEqual(text_frame._txBody.bodyPr.normAutofit.fontScale, 50.0)

    @patch('app.translate.to_translate.time.sleep')
    @patch('app.translate.to_translate._translate_openai')
    def test_style_failure_uses_dominant_style_fallback(self, translate_openai, sleep):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        paragraph.clear()
        paragraph.add_run().text = 'Long dominant text '
        short = paragraph.add_run()
        short.text = 'x'
        short.font.italic = True
        plan = build_inline_plan(paragraph)
        item = {'text': plan.payload, **plan.translation_metadata()}
        translate_openai.side_effect = [
            'broken', 'broken', 'broken', 'Texte traduit complet',
        ]

        result = to_translate._translate_text_block(
            {'id': 1, 'server': 'openai', 'model': 'model'}, item
        )
        self.assertTrue(result['fallback_used'])
        plan.apply(result['translated_text'], fallback_used=True)
        self.assertEqual(paragraph.text, 'Texte traduit complet')
        self.assertEqual(paragraph.runs[0].text, 'Texte traduit complet')
        self.assertEqual(paragraph.runs[1].text, '')

    @patch('app.translate.to_translate.time.sleep')
    @patch('app.translate.to_translate._translate_openai', return_value='broken')
    def test_structural_anchor_failure_never_uses_style_fallback(self, translate_openai, sleep):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
        paragraph.clear()
        linked = paragraph.add_run()
        linked.text = 'Documentation'
        linked.hyperlink.address = 'https://example.com'
        plan = build_inline_plan(paragraph)
        item = {'text': plan.payload, **plan.translation_metadata()}

        with self.assertRaises(to_translate.FatalError):
            to_translate._translate_text_block(
                {'id': 2, 'server': 'openai', 'model': 'model'}, item
            )


class PptxHighFidelityTests(unittest.TestCase):
    def _run_translation(self, source: Path, target: Path, mode: str):
        with patch.object(
            powerpoint.to_translate, 'translate_batch', _fake_translate_batch
        ), patch.object(
            powerpoint.to_translate, 'complete', return_value=True
        ), patch.object(powerpoint.to_translate, 'error'):
            return powerpoint.start({
                'id': 9,
                'file_path': str(source),
                'target_file': str(target),
                'type': mode,
                'lang': '法语',
            })

    def test_translation_preserves_geometry_styles_tables_and_relationships(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            image = directory / 'pixel.png'
            image.write_bytes(PNG_1X1)
            source = directory / 'source.pptx'
            target = directory / 'target.pptx'
            _build_complex_presentation(source, image)
            source_presentation = Presentation(source)
            baseline = powerpoint._capture_presentation_baseline(source_presentation)

            self.assertTrue(self._run_translation(source, target, 'trans_all_only_inherit'))
            output = Presentation(target)
            snapshot = powerpoint._snapshot_slide(output.slides[0])
            self.assertEqual(snapshot['structure'], baseline['slides'][0]['structure'])
            self.assertEqual(snapshot['relationships'], baseline['slides'][0]['relationships'])

            texts = _all_text(output.slides[0])
            self.assertIn('Bonjour monde', texts)
            self.assertIn('Premier paragraphe', texts)
            self.assertIn('Deuxieme paragraphe', texts)
            translated_first = next(
                shape for shape in output.slides[0].shapes
                if shape.has_text_frame and shape.text_frame.text == 'Bonjour monde'
            )
            self.assertTrue(translated_first.text_frame.paragraphs[0].runs[0].font.bold)
            self.assertTrue(translated_first.text_frame.paragraphs[0].runs[1].font.italic)

            shape_types = [shape.shape_type for shape in output.slides[0].shapes]
            self.assertIn(MSO_SHAPE_TYPE.PICTURE, shape_types)
            self.assertIn(MSO_SHAPE_TYPE.CHART, shape_types)
            chart_shape = next(
                shape for shape in output.slides[0].shapes
                if shape.shape_type == MSO_SHAPE_TYPE.CHART
            )
            self.assertEqual(chart_shape.chart.series[0].name, 'Series')
            hyperlink_run = next(
                run for shape in output.slides[0].shapes if shape.has_text_frame
                for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                if run.text == 'Documentation francaise'
            )
            self.assertEqual(hyperlink_run.hyperlink.address, 'https://example.com/docs')
            long_text_shape = next(
                shape for shape in output.slides[0].shapes
                if shape.has_text_frame
                and shape.text_frame.text.startswith('Deuxieme forme')
            )
            long_run = long_text_shape.text_frame.paragraphs[0].runs[0]
            self.assertEqual(long_text_shape.width, Inches(3))
            self.assertEqual(long_text_shape.height, Inches(0.8))
            self.assertEqual(long_run.font.size.pt, 18)
            self.assertEqual(
                long_text_shape.text_frame.auto_size,
                MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
            )
            font_scale = long_text_shape.text_frame._txBody.bodyPr.normAutofit.fontScale
            self.assertGreaterEqual(font_scale, 60.0)
            self.assertLess(font_scale, 100.0)

    def test_bilingual_mode_keeps_original_and_uses_stable_shape_ids(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            image = directory / 'pixel.png'
            image.write_bytes(PNG_1X1)
            source = directory / 'source.pptx'
            target = directory / 'target.pptx'
            _build_complex_presentation(source, image)
            source_presentation = Presentation(source)
            source_snapshot = powerpoint._snapshot_slide(source_presentation.slides[0])

            self.assertTrue(self._run_translation(source, target, 'trans_all_both_inherit'))
            output = Presentation(target)
            self.assertEqual(len(output.slides), 2)
            original_snapshot = powerpoint._snapshot_slide(output.slides[0])
            translated_snapshot = powerpoint._snapshot_slide(output.slides[1])
            self.assertEqual(original_snapshot, source_snapshot)
            self.assertEqual(translated_snapshot['structure'], source_snapshot['structure'])
            self.assertEqual(translated_snapshot['relationships'], source_snapshot['relationships'])
            self.assertIn('Hello world', _all_text(output.slides[0]))
            self.assertIn('Bonjour monde', _all_text(output.slides[1]))
            original_long_shape = next(
                shape for shape in output.slides[0].shapes
                if shape.has_text_frame and shape.text_frame.text == 'Second shape'
            )
            translated_long_shape = next(
                shape for shape in output.slides[1].shapes
                if shape.has_text_frame
                and shape.text_frame.text.startswith('Deuxieme forme')
            )
            self.assertEqual(
                original_long_shape.text_frame.auto_size,
                MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT,
            )
            self.assertEqual(
                translated_long_shape.text_frame.auto_size,
                MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
            )

    def test_bilingual_mode_preserves_multi_slide_order(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            directory = Path(temp_dir)
            source = directory / 'source.pptx'
            target = directory / 'target.pptx'
            presentation = Presentation()
            for text in ('Alpha page', 'Beta page'):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = text
            presentation.save(source)

            def translate_pages(trans, texts, event):
                for item in texts:
                    item['text'] = item['text'].replace('Alpha page', 'Page alpha').replace(
                        'Beta page', 'Page beta'
                    )
                    item['count'] = 1
                    item['complete'] = True
                    item['fallback_used'] = False
                return True

            with patch.object(
                powerpoint.to_translate, 'translate_batch', translate_pages
            ), patch.object(
                powerpoint.to_translate, 'complete', return_value=True
            ), patch.object(powerpoint.to_translate, 'error'):
                self.assertTrue(powerpoint.start({
                    'id': 10,
                    'file_path': str(source),
                    'target_file': str(target),
                    'type': 'trans_all_both_inherit',
                    'lang': '法语',
                }))

            output = Presentation(target)
            self.assertEqual(
                [_all_text(slide)[0] for slide in output.slides],
                ['Alpha page', 'Page alpha', 'Beta page', 'Page beta'],
            )

    def test_invalid_saved_presentation_is_removed(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            target = Path(temp_dir) / 'broken.pptx'

            def write_broken(path):
                Path(path).write_bytes(b'not a pptx')

            fake_presentation = SimpleNamespace(save=write_broken)
            baseline = {'slide_width': 1, 'slide_height': 1, 'slides': []}
            with patch('app.translate.powerpoint.Presentation', side_effect=ValueError('broken')):
                with self.assertRaisesRegex(ValueError, 'broken'):
                    powerpoint._save_and_validate_presentation(
                        fake_presentation, str(target), baseline, False
                    )
            self.assertFalse(target.exists())


if __name__ == '__main__':
    unittest.main()
